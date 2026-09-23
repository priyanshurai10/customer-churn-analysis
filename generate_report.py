import docx

# Create a new Document
doc = docx.Document()

# Add a title
doc.add_heading('Project Report: Customer Churn Analysis', 0)

# Introduction
doc.add_heading('1. Introduction', level=1)
doc.add_paragraph(
    "Hi, I am a B.Tech CSE student. This is my data analytics project created as part of the IBM SkillsBuild / BharatCares internship. "
    "The project focuses on Customer Churn Analysis, which means studying the data of customers who stop doing business with a company."
)

# Problem Statement
doc.add_heading('2. Problem Statement', level=1)
doc.add_paragraph(
    "The main goal of this project is to analyze customer data to understand:\n"
    "- How many customers are leaving.\n"
    "- Which customer groups have a higher churn rate.\n"
    "- Whether factors like contract type, internet service, and customer tenure affect churn.\n"
    "- What possible business actions the company can take to retain customers."
)

# Dataset
doc.add_heading('3. Dataset', level=1)
doc.add_paragraph(
    "I used the IBM Telco Customer Churn sample dataset. It includes information about telecom customers, such as "
    "their tenure, monthly charges, contract type, internet service, and whether they left the company (churned)."
)

# Data Cleaning
doc.add_heading('4. Data Cleaning', level=1)
doc.add_paragraph(
    "Before analyzing the data, I cleaned it using Pandas in Python. I:\n"
    "- Loaded the dataset from the CSV link.\n"
    "- Converted the 'TotalCharges' column into a numeric format (it had some empty spaces).\n"
    "- Removed a few rows that had missing values in TotalCharges.\n"
    "- Created a simple 'ChurnFlag' column (1 for Yes, 0 for No) to make calculations easier.\n"
    "- Created tenure groups like '0-1 Year' to make the charts simpler to read."
)

# Analysis
doc.add_heading('5. Analysis', level=1)
doc.add_paragraph(
    "I used Matplotlib to create simple charts and found some clear patterns:\n"
    "- Churn by Contract: Customers on a 'Month-to-Month' contract have a much higher churn rate compared to those on 1-year or 2-year contracts.\n"
    "- Churn by Internet Service: Customers using 'Fiber optic' internet leave more often than those using DSL or no internet.\n"
    "- Churn by Tenure: Most churn happens in the first year (0-1 Year group)."
)

# Machine Learning Model
doc.add_heading('6. Machine Learning Model', level=1)
doc.add_paragraph(
    "I built a basic Logistic Regression model using Scikit-Learn to see if I could predict churn. "
    "I split the data into training (80%) and testing (20%) sets. I converted text columns into numbers using one-hot encoding "
    "and scaled the numerical columns so the model could learn better."
)

# Results
doc.add_heading('7. Results', level=1)
doc.add_paragraph(
    "When I tested the model on the test data, it gave an ROC-AUC score of around 0.84 and an F1 score of around 0.59. "
    "This means the model is decent at recognizing risk patterns, but it is just a basic baseline and does not guarantee exactly who will leave."
)

# Risk / Opportunity / Possible Action
doc.add_heading('8. Risk, Opportunity, and Possible Action', level=1)
doc.add_paragraph(
    "Risk:\n"
    "Customers with Month-to-Month contracts and Fiber Optic internet are at high risk of leaving.\n\n"
    "Opportunity:\n"
    "Customers on 1-Year or 2-Year contracts are very loyal. The company should focus on getting more people into these contracts.\n\n"
    "Possible Action:\n"
    "- Better Onboarding: Give more support during the first 6-12 months.\n"
    "- Targeted Retention Offers: Offer a small discount to month-to-month customers if they switch to a 1-year contract.\n"
    "- Check the Fiber Optic service to see if there are technical issues causing people to leave."
)

# Technologies Used
doc.add_heading('9. Technologies Used', level=1)
doc.add_paragraph(
    "- Python\n"
    "- Pandas & NumPy (Data Cleaning and Handling)\n"
    "- Matplotlib (Data Visualization)\n"
    "- Scikit-Learn (Basic Machine Learning)\n"
    "- Streamlit (Dashboard UI)"
)

# Conclusion
doc.add_heading('10. Conclusion', level=1)
doc.add_paragraph(
    "In conclusion, this project helped me learn how to take raw data, clean it, find useful patterns, "
    "and suggest basic business actions. It was a great learning experience for practical data analytics."
)

# Save the document
doc.save('c:/Users/priya/Music/ineternship/customer-churn-analysis/project_report.docx')

# Also save as markdown for a text version
with open('c:/Users/priya/Music/ineternship/customer-churn-analysis/project_report.md', 'w') as f:
    f.write("# Project Report: Customer Churn Analysis\n\n(See project_report.docx for the full formatted report)")
